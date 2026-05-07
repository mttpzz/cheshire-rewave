from cat.logs.cat_logger import get_plugin_logger     # type: ignore
from cat.mad_hatter.decorators import tool  # type: ignore
import os
import json
from datetime import datetime
from fpdf import FPDF
from pypdf import PdfReader
from docx import Document
import html


# --- LOGGER --------------------------------------------------------------------------------------------------------------------
# start the logger with its plugin name
log = get_plugin_logger("file-manager")


# --- ENV VARIABLES -------------------------------------------------------------------------------------------------------------
BASE_FOLDER_CAT = os.getenv('BASE_FOLDER_CAT')
BASE_FOLDER_USER = os.getenv('BASE_FOLDER_USER')


# --- PATH SECURITY FUNCTION ----------------------------------------------------------------------------------------------------
# Security and utility function to manage user paths safely (Path traversal prevention)
def get_user_path(cat, folder, filename=None):
    user_id = cat.user_id
    base_path = os.path.join(folder, user_id)
    
    os.makedirs(base_path, exist_ok=True)
        
    if filename:
        # os.path.basename impedisce attacchi di tipo path traversal
        return os.path.join(base_path, os.path.basename(filename))
    
    return base_path


# --- TOOL: FILE WRITE ----------------------------------------------------------------------------------------------------------
@tool(return_direct=True, examples=["Creami un pdf con il riassunto del documento che ti ho inviato.", "Voglio un file Word che riassuma la nostra conversazione."])
def create_file(input_json, cat):
    """
    This tool create or write a file in a specific folder for the user and return the folder in which the file is saved.
    The input MUST be a Python dictionary which contains the text to be included in the file, the format of the file (word, pdf or txt) and the name of the file (optional, if not provided it will be generated with a timestamp),
    for example: {"text": "Questo è il testo da inserire nel documento.", "format": "pdf", "filename": "riassunto"}
    """
    log.info("Starting file writing tool.")

    # parsing input
    try:
        data = json.loads(input_json) if isinstance(input_json, str) else input_json
        file_text = data.get("text", "")
        file_format = (data.get("format") or "pdf").lower().strip()
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        file_name = data.get("filename") or f"doc_{timestamp}"
    except Exception as e:
        log.error(f"❌ Error while parsing input: {str(e)}.")
        return "❌ Errore nel formato dei dati. Riprova."
    
    # file creation
    try:
        if file_format == "pdf":
            try:
                file_text.encode("cp1252")
                file_name += ".pdf"
                full_path = get_user_path(cat, BASE_FOLDER_CAT, file_name)
                pdf = FPDF()
                pdf.add_page()
                pdf.set_font("Arial", size=12)
                pdf.multi_cell(0, 10, text=file_text)
                pdf.output(full_path)
            except UnicodeEncodeError:
                # fallback to DOCX (Unicode-safe) if the text contains characters that cannot be encoded in cp1252 (e.g., emojis, CJK, cyrillic)
                log.warning("⚠️ Non-latin1 chars detected, switching PDF to DOCX.")
                file_name += ".docx"
                full_path = get_user_path(cat, BASE_FOLDER_CAT, file_name)

                doc = Document()
                for line in file_text.split("\n"):
                    doc.add_paragraph(line)
                doc.save(full_path)
        elif file_format in ["word", "docx"]:
            file_name += ".docx"
            full_path = get_user_path(cat, BASE_FOLDER_CAT, file_name)

            doc = Document()
            for line in file_text.split("\n"):
                doc.add_paragraph(line)
            doc.save(full_path)
        else:   # txt files
            file_name += ".txt"
            full_path = get_user_path(cat, BASE_FOLDER_CAT, file_name)

            with open(full_path, "w", encoding="utf-8") as f:
                f.write(file_text)
    
    except Exception as e:
        log.error(f"❌ Error during file creation: {str(e)}.")
        return "❌ Errore durante la creazione del file. Riprova."

    # folder for file access
    user_path = get_user_path(cat, BASE_FOLDER_USER)
    user_path_link = f'<a href="{html.escape(user_path, quote=True)}" target="_blank">{html.escape(user_path)}</a>'
    
    return f"✅ Fatto! Puoi vedere il file **{file_name}** in questa cartella: {user_path_link}"


# --- TOOL: FILE LIST -----------------------------------------------------------------------------------------------------------
@tool(return_direct=True, examples=["Quali file ho nella mia cartella?", "Mostrami i miei documenti."])
def list_files(_, cat):
    """
    List all files in the user's folder.
    The input is an empty object {}.
    """
    log.info("Starting files listing tool.")

    user_path = get_user_path(cat, BASE_FOLDER_CAT)
    try:
        files = os.listdir(user_path)
        if not files:
            return "La tua cartella è attualmente vuota."
        
        files_list = "\n- ".join(files)
        return f"Ecco i file nella tua cartella:\n- {files_list}"
    
    except Exception as e:
        log.error(f"❌ Error while listing the files in the folder: {str(e)}.")
        return "❌ Errore durante la lettura della cartella. Riprova."


# --- TOOL: FILE READ -----------------------------------------------------------------------------------------------------------
@tool(return_direct=True, examples=["Leggimi il contenuto di 'note.pdf'", "Cosa c'è scritto in 'documento.docx'?"])
def read_file(input_json, cat):
    """
    Read the content of a specific file (txt, pdf or docx) in the user folder.
    Input MUST be a JSON with the key 'filename',
    for example: {"filename": "note.docx"}
    """
    log.info("Starting file reading tool.")

    try:
        data = json.loads(input_json) if isinstance(input_json, str) else input_json
        filename = data.get("filename")
        if not filename:
            return "❌ Nome file mancante. Specifica il file da leggere."
        path = get_user_path(cat, BASE_FOLDER_CAT, filename)

        if os.path.exists(path):
            ext = filename.lower().split('.')[-1]   # file extension
            text_read = ""
            
            if ext == "pdf":
                reader = PdfReader(path)
                for page in reader.pages:
                    text_read += (page.extract_text() or "") + "\n"
            elif ext == "docx":
                doc = Document(path)
                text_read = "\n".join([para.text for para in doc.paragraphs])
            elif ext == "doc":
                return "⚠️ Il formato '.doc' non è supportato. Convertilo in '.docx' prima di leggerlo."
            else:            
                with open(path, "r", encoding="utf-8") as f:
                    text_read = f.read()
            
            if not text_read.strip():
                return f"Il file '{filename}' sembra essere vuoto o non leggibile."
            else:
                return f"Contenuto di '{filename}':\n\n{text_read}"
        
        return f"❌ Il file '{filename}' non esiste."

    except Exception as e:
        log.error(f"❌ Error while reading the file: {str(e)}")
        return "❌ Errore durante la lettura del file. Riprova."


# --- TOOL: FILE RENAME ---------------------------------------------------------------------------------------------------------
@tool(return_direct=True, examples=["Rinomina 'vecchio.pdf' in 'nuovo.pdf'", "Voglio che 'appunti.docx' si chiami 'riassunto.docx'."])
def rename_file(input_json, cat):
    """
    Rename a file in the user's folder.
    Input MUST be a JSON with 'old_name' and 'new_name',
    for example: {"old_name": "vecchio.txt", "new_name": "nuovo.txt"}
    """
    log.info("Starting file renaming tool.")

    old_name = None
    try:
        data = json.loads(input_json) if isinstance(input_json, str) else input_json
        old_name = data.get("old_name")
        new_name = data.get("new_name")
        if not old_name or not new_name:
            return "❌ Nomi file mancanti. Specifica nome vecchio e nuovo."
        old_path = get_user_path(cat, BASE_FOLDER_CAT, old_name)
        new_path = get_user_path(cat, BASE_FOLDER_CAT, new_name)

        if old_path == new_path:
            return f"⚠️ Il nome '{new_name}' è già quello attuale."
        
        if os.path.exists(old_path):
            os.rename(old_path, new_path)
            return f"✅ File rinominato con successo in '{new_name}'."
        
        return f"❌ Il file '{old_name}' non esiste."
    
    except Exception as e:
        log.error(f"❌ Error while renaming the file '{old_name}': {str(e)}")
        return "❌ Errore durante la rinomina del file. Riprova."


# --- TOOL: FILE DELETE ---------------------------------------------------------------------------------------------------------
@tool(return_direct=True, examples=["Elimina 'da_cancellare.pdf'", "Voglio cancellare 'vecchio.pdf'."])
def delete_file(input_json, cat):
    """
    Delete permanently a file from the user's folder.
    Input MUST be a JSON with the key 'filename',
    for example: {"filename": "da_cancellare.txt"}
    """
    log.info("Starting file deleting tool.")

    filename = None
    try:
        data = json.loads(input_json) if isinstance(input_json, str) else input_json
        filename = data.get("filename")
        if not filename:
            return "❌ Nome file mancante. Specifica il file da eliminare."
        path = get_user_path(cat, BASE_FOLDER_CAT, filename)
        
        if os.path.exists(path):
            os.remove(path)
            return f"✅ File '{filename}' eliminato con successo."

        return f"❌ Il file '{filename}' non è stato trovato."
    
    except Exception as e:
        log.error(f"❌ Error while deleting the file '{filename}': {str(e)}")
        return "❌ Errore durante l'eliminazione del file."
